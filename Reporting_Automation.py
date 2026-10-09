"""Автоматизация управленческого отчёта P&L из ОСВ, выгруженных из 1С.
 
Как работает программа (по шагам):
1. Берёт одну ОСВ из папки data_input.
2. Вытаскивает из неё статьи по счетам (6010, 6100, ... 7400).
3. Присоединяет справочник соответствий mapping.xlsx (счёт + статья -> строка P&L).
4. Применяет доли газа/логистики из settings.xlsx.
5. Складывает по строкам P&L, считает прибыли, сохраняет Excel в data_output.
6. Проверяет, что разнесённое равно итогу ОСВ.
"""
import re
from pathlib import Path
 
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
 
pd.set_option("display.max_columns", None)
pd.set_option("display.width", 250)
pd.options.display.float_format = "{:,.2f}".format
 
# ---------------------------------------------------------------- настройки
INPUT_DIR = Path("data_input")
OUTPUT_DIR = Path("data_output")
MAPPING_FILE = Path("mapping.xlsx")
LAYOUT_FILE = Path("layout.xlsx")
SETTINGS_FILE = Path("settings.xlsx")
 
MONTHS = {
    "Январь": 1, "Февраль": 2, "Март": 3, "Апрель": 4,
    "Май": 5, "Июнь": 6, "Июль": 7, "Август": 8,
    "Сентябрь": 9, "Октябрь": 10, "Ноябрь": 11, "Декабрь": 12,
}
 
# Запасное правило: куда отнести статью, которой нет в справочнике.
# (счёт, раздел, строка P&L, знак)
FALLBACK = [
    ("6010", "Доход", "Доход - прочее", 1),
    ("6200", "Доход", "Доход - прочее", 1),
    ("7010", "Себестоимость", "Прочие", 1),
    ("7110", "Расходы по реализации", "Прочее", 1),
    ("7210", "Административные расходы", "Прочее - услуги", 1),
    ("6100", "Прочее", "Прочее", -1),
    ("7300", "Прочее", "Финансовые расходы", 1),
    ("7400", "Прочее", "Курсовые разницы", 1),
]
 
# Правило по счёту целиком: любая статья этого счёта идёт в одну строку P&L,
# в справочнике названия не нужны, и это не считается «нераспознанной» статьёй.
# (счёт: (раздел, строка P&L, знак))
ACCOUNT_RULES = {
    "7710": ("Прочее", "КПН", 1),
}
 
 
# ------------------------------------------------- 1. чтение ОСВ из файла
def find_input_file():
    """Возвращает путь к единственной ОСВ в папке data_input."""
    files = [f for f in INPUT_DIR.glob("*.xlsx") if not f.name.startswith("~$")]
    if len(files) != 1:
        raise ValueError(
            f"В папке {INPUT_DIR} должен лежать ровно один файл ОСВ, найдено: {len(files)}"
        )
    return files[0]
 
 
def clean_block(block, account):
    """Из блока одного счёта оставляет только строки-статьи."""
    block = block[[0, 4, 5, "indent"]].copy()
    block.columns = ["article", "debit", "credit", "indent"]
    block["article"] = block["article"].astype(str).str.strip()
    block["debit"] = pd.to_numeric(block["debit"], errors="coerce")
    block["credit"] = pd.to_numeric(block["credit"], errors="coerce")
 
    rows = block[(block["indent"] > 0) & (block["article"] != "Головное подразделение")]
    level = rows["indent"].min()
    rows = rows[rows["indent"] == level]
    rows = rows[rows["debit"].notna() | rows["credit"].notna()]
 
    rows = rows.assign(account=account)
    return rows[["account", "article", "debit", "credit"]]
 
 
def get_period(df):
    """Достаёт период из заголовка «... за Октябрь 2024 г.» -> «2024-10»."""
    first_col = df[0].astype(str)
    found = first_col[first_col.str.contains(r"за \S+ \d{4}")].iloc[0]
    month_name, year = re.search(r"за (\S+) (\d{4})", found).groups()
    return f"{year}-{MONTHS[month_name]:02d}"
 
 
def load_osv(path):
    """Читает файл ОСВ и возвращает таблицу: счёт, статья, сумма, период."""
    df = pd.read_excel(path, sheet_name="ОСВ", header=None)
 
    ws = load_workbook(path)["ОСВ"]
    df["indent"] = [ws.cell(row=i + 1, column=1).alignment.indent for i in range(len(df))]
 
    first_col = df[0].astype(str)
    headers = first_col[first_col.str.contains("по счету")]
    accounts = headers.str.extract(r"по счету (\d+)")[0]
 
    starts = headers.index.tolist()
    ends = starts[1:] + [len(df)]
 
    tables = []
    for account, start, end in zip(accounts, starts, ends):
        tables.append(clean_block(df.iloc[start:end], account))
 
    osv = pd.concat(tables, ignore_index=True)
    # доходы (6xxx) берём по кредиту, расходы (7xxx) по дебету
    osv["amount"] = osv["credit"].where(osv["account"].str.startswith("6"), osv["debit"])
    osv["amount"] = osv["amount"].fillna(0)
    osv["period"] = get_period(df)
    return osv
 
 
# ------------------------------------------------- 2. справочник соответствий
def apply_mapping(osv):
    """Присоединяет к каждой статье строку P&L; нераспознанные идут по запасному правилу."""
    mapping = pd.read_excel(MAPPING_FILE, dtype={"account": str})
    mapping["article"] = mapping["article"].str.strip()
    merged = osv.merge(mapping, on=["account", "article"], how="left")
 
    for account, (section, pl_line, sign) in ACCOUNT_RULES.items():
        mask = (merged["account"] == account) & merged["pl_line"].isna()
        merged.loc[mask, "section"] = section
        merged.loc[mask, "pl_line"] = pl_line
        merged.loc[mask, "sign"] = sign
 
    fallback = pd.DataFrame(FALLBACK, columns=["account", "fb_section", "fb_pl_line", "fb_sign"])
    merged = merged.merge(fallback, on="account", how="left")
    merged["recognized"] = merged["pl_line"].notna()
 
    merged["section"] = merged["section"].fillna(merged["fb_section"])
    merged["pl_line"] = merged["pl_line"].fillna(merged["fb_pl_line"])
    merged["sign"] = merged["sign"].fillna(merged["fb_sign"])
    return merged
 
 
def get_unrecognized(merged):
    """Статьи, которых нет в справочнике, но с ненулевой суммой."""
    mask = (~merged["recognized"]) & (merged["amount"] != 0)
    return merged[mask][["account", "article", "amount", "pl_line"]]
 
 
# ------------------------------------------------- 3. настройки (доли 60/40)
def load_settings():
    table = pd.read_excel(SETTINGS_FILE)
    return dict(zip(table["key"], table["value"]))
 
 
def validate_settings(settings):
    """Останавливает программу, если доли заданы неверно."""
    for key in ("gas_share", "logistics_share"):
        if key not in settings:
            raise ValueError(f"В {SETTINGS_FILE} нет строки {key}")
        if not 0 <= settings[key] <= 1:
            raise ValueError(f"{key} должна быть числом от 0 до 1, сейчас: {settings[key]}")
    total = settings["gas_share"] + settings["logistics_share"]
    if abs(total - 1) > 1e-9:
        raise ValueError(f"Доли газа и логистики в сумме должны давать 100%, сейчас: {total:.0%}")
 
 
def ask_settings_confirmation():
    """Показывает доли, ждёт Enter, читает файл заново и проверяет."""
    settings = load_settings()
    print()
    print(f"Значения из {SETTINGS_FILE.resolve()}:")
    print(f"  Доля СУГ на ТМЗ - основное производство: {settings['gas_share']:.0%}")
    print(f"  Доля СУГ на ТМЗ - логистические расходы: {settings['logistics_share']:.0%}")
    print()
    print("Если изменений нет, нажмите Enter, чтобы продолжить.")
    print("Если нужно изменить, поправьте файл, сохраните его (Ctrl+S) и нажмите Enter.")
    input()
 
    settings = load_settings()  # читаем заново: вдруг файл правили
    validate_settings(settings)
    return settings
 
 
def apply_shares(merged, settings):
    """Считает итоговую сумму: сумма * знак * доля."""
    share_by_line = {
        "ТМЗ - основное производство": settings["gas_share"],
        "ТМЗ - логистические расходы": settings["logistics_share"],
    }
    merged["share"] = merged["pl_line"].map(share_by_line).fillna(1)
    merged["final_amount"] = merged["amount"] * merged["sign"] * merged["share"]
    return merged
 
 
# ------------------------------------------------- 4. свод и показатели
def build_report(merged):
    """Складывает по строкам P&L, переводит в тыс. тенге."""
    report = merged.groupby(["section", "pl_line"], as_index=False)["final_amount"].sum()
    report["final_amount"] = report["final_amount"] / 1000
    return report
 
 
def calc_results(report):
    """Считает итоги разделов и прибыли по формулам старого отчёта."""
 
    def line(section, pl_line):
        mask = (report["section"] == section) & (report["pl_line"] == pl_line)
        return report.loc[mask, "final_amount"].sum()
 
    totals = report.groupby("section")["final_amount"].sum()
    income = totals.get("Доход", 0)
    cost = totals.get("Себестоимость", 0)
    selling = totals.get("Расходы по реализации", 0)
    admin = totals.get("Административные расходы", 0)
 
    gross_profit = income - cost
    ebitda = gross_profit - selling - admin
 
    depreciation = line("Прочее", "Амортизация - Амортизация") + line(
        "Прочее", "Амортизация - Амортизация - офис"
    )
    operating_profit = ebitda - depreciation
 
    other_result = (
        line("Прочее", "Курсовые разницы")
        + line("Прочее", "Выбытие активов")
        + line("Прочее", "Прочее")
    )
    ebit = operating_profit - other_result
 
    finance = line("Прочее", "Финансовые расходы") + line("Прочее", "Начисленные финансовые расходы")
    ebt = ebit - finance
 
    tax = line("Прочее", "КПН")
    net_profit = ebt - tax
 
    return {
        "income": income, "cost": cost, "selling": selling, "admin": admin,
        "gross_profit": gross_profit, "ebitda": ebitda,
        "operating_profit": operating_profit, "ebit": ebit,
        "ebt": ebt, "net_profit": net_profit,
    }
 
 
# ------------------------------------------------- 5. таблица отчёта и Excel
def build_pl(report, results):
    """Собирает итоговую таблицу в порядке отчёта (строки из layout.xlsx)."""
    layout = pd.read_excel(LAYOUT_FILE)
    full = layout.merge(report, on=["section", "pl_line"], how="left")
    full["final_amount"] = full["final_amount"].fillna(0)
    block_sum = full.groupby("block")["final_amount"].sum()
 
    rows = []
 
    def add(label, value, kind):
        rows.append({"label": label, "value": value, "kind": kind})
 
    def add_block(title, block, header=True):
        if header:
            add(title, block_sum[block], "subtotal")
        part = full[full["block"] == block]
        for name, value in zip(part["pl_line"], part["final_amount"]):
            add(name, value, "line")
 
    add_block("Доход", "Доход")
 
    add("Себестоимость - итого", results["cost"], "total")
    add_block("Себестоимость - общие", "Себестоимость - общие")
    add_block("Себестоимость - специфические", "Себестоимость - специфические")
    add("Валовая прибыль", results["gross_profit"], "result")
 
    add("Расходы по реализации - итого", results["selling"], "total")
    add_block("Расходы по реализации - общие", "Расходы по реализации - общие")
    add_block("Расходы по реализации - специфические", "Расходы по реализации - специфические")
 
    add("Административные расходы - итого", results["admin"], "total")
    add_block("Административные расходы - общие", "Административные расходы - общие")
    add("EBITDA", results["ebitda"], "result")
 
    add_block("Амортизация", "Амортизация")
    add("Операционная прибыль", results["operating_profit"], "result")
 
    add_block("Прочие прибыли / убытки", "Прочие прибыли / убытки")
    add("EBIT", results["ebit"], "result")
 
    add_block("Финансовые расходы", "Финансовые расходы", header=False)
    add("EBT", results["ebt"], "result")
 
    add_block("КПН", "Налог", header=False)
    add("Чистая прибыль", results["net_profit"], "result")
 
    pl = pd.DataFrame(rows)
    pl["value"] = pl["value"].round(2)
    return pl
 
 
def save_excel(pl, out_path):
    """Сохраняет отчёт в Excel и оформляет его."""
    OUTPUT_DIR.mkdir(exist_ok=True)
    pl_out = pl[["label", "value"]].rename(columns={"label": "Статья", "value": "тыс. тенге"})
    pl_out.to_excel(out_path, index=False, sheet_name="PL")
 
    wb = load_workbook(out_path)
    ws = wb["PL"]
 
    bold = Font(bold=True)
    gray = PatternFill("solid", fgColor="D9D9D9")
 
    for cell in ws[1]:
        cell.font = bold
 
    for row_num, kind in enumerate(pl["kind"], start=2):
        label_cell = ws.cell(row=row_num, column=1)
        value_cell = ws.cell(row=row_num, column=2)
 
        value_cell.number_format = '#,##0.00;-#,##0.00;"-"'
 
        if kind == "line":
            label_cell.alignment = Alignment(indent=2)
        else:
            label_cell.font = bold
            value_cell.font = bold
 
        if kind == "result":
            label_cell.fill = gray
            value_cell.fill = gray
 
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 16
    ws.freeze_panes = "A2"
 
    wb.save(out_path)
 
 
# ------------------------------------------------- 6. контроль
def check_totals(osv, merged):
    """Сравнивает сумму по каждому счёту в ОСВ и в отчёте."""
    placed = merged[merged["pl_line"].notna()]
    check = pd.DataFrame(
        {
            "osv": osv.groupby("account")["amount"].sum(),
            "report": (placed["amount"] * placed["share"]).groupby(placed["account"]).sum(),
        }
    )
    check["diff"] = (check["osv"] - check["report"]).round(2)
    return check
 
 
# ------------------------------------------------- запуск
def main():
    try:
        file = find_input_file()
        osv = load_osv(file)
        period = osv["period"].iloc[0]
        print(f"Файл: {file.name} | период: {period}")
 
        settings = ask_settings_confirmation()
    except ValueError as error:
        print("ОШИБКА:", error)
        return
 
    merged = apply_mapping(osv)
    merged = apply_shares(merged, settings)
 
    print("\nНераспознанные статьи с суммой:")
    print(get_unrecognized(merged))
 
    report = build_report(merged)
    results = calc_results(report)
    pl = build_pl(report, results)
 
    out_path = OUTPUT_DIR / f"PL_{period}.xlsx"
    save_excel(pl, out_path)
    print("\nСохранено:", out_path)
    for name in ("gross_profit", "ebitda", "operating_profit", "ebit", "ebt", "net_profit"):
        print(f"  {name}: {results[name]:,.2f}")
 
    check = check_totals(osv, merged)
    print("\nКонтроль по счетам:")
    print(check)
    if (check["diff"] != 0).any():
        print("ВНИМАНИЕ: часть сумм ОСВ не попала в отчёт!")
    else:
        print("Контроль пройден: всё разнесённое равно итогу ОСВ.")
 
 
if __name__ == "__main__":
    main()
    